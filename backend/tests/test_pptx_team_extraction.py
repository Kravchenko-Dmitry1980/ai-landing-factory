"""PPTX team extraction: slides, tables, grouped shapes, false positive guard."""

from __future__ import annotations

from pathlib import Path
from uuid import uuid4

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.models.domain import utc_now
from app.schemas.extraction import ExtractionPayload, ExtractionResult, FileExtraction
from app.services.analysis.contract_builder import ContractBuilderService
from app.services.contract_fidelity.completeness_gate import ContractCompletenessGate
from app.services.contract_fidelity.landing_document_detector import LandingDocumentDetector
from app.services.contract_fidelity.presentation_landing_synthesizer import (
    PresentationLandingSynthesizer,
)
from app.services.contract_fidelity.source_type_detector import SourceTypeDetector
from app.services.contract_fidelity.structured_landing_parser import StructuredLandingParser
from app.services.evidence.multi_source_assembler import MultiSourceEvidenceAssembler
from app.services.evidence.people_extractor import extract_people_from_text
from app.services.evidence.source_inventory import SourceInventoryBuilder
from app.services.extraction.pptx_extractor import PptxExtractor
from app.services.export.styled_html_exporter import ExportTheme, StyledHtmlExporter
from app.config import settings
from app.repositories.contract_repository import ContractRepository

INDLAB_PRES = (
    Path(__file__).parents[2]
    / "test_corpus"
    / "golden"
    / "indlab_telegram_news"
    / "sources"
    / "01_presentation.pptx.txt"
)
KSK_PRES = (
    Path(__file__).parents[2]
    / "test_corpus"
    / "golden"
    / "ksk_it_barrier"
    / "sources"
    / "01_presentation.pptx.txt"
)


def _make_builder() -> ContractBuilderService:
    builder = ContractBuilderService.__new__(ContractBuilderService)
    builder._detector = LandingDocumentDetector()
    builder._source_type_detector = SourceTypeDetector()
    builder._structured_parser = StructuredLandingParser()
    builder._presentation_synthesizer = PresentationLandingSynthesizer()
    builder._completeness_gate = ContractCompletenessGate()
    builder._inventory_builder = SourceInventoryBuilder()
    builder._multi_source_assembler = MultiSourceEvidenceAssembler()
    return builder


def _pptx_only_extraction(text: str, filename: str = "deck.pptx") -> ExtractionResult:
    return ExtractionResult(
        project_id=uuid4(),
        payload=ExtractionPayload(),
        files=[
            FileExtraction(
                filename=filename,
                file_type="pptx",
                extracted_text=text,
                metadata={"slides_count": text.count("Slide ")},
            )
        ],
        extracted_at=utc_now(),
    )


def test_team_slide_with_fio_extracted() -> None:
    text = """
Slide 1:
Команда проекта
Кравченко Дмитрий — тимлид
Ерюкова Татьяна — помощник тимлида
"""
    people = extract_people_from_text(text, in_team_section=True)
    names = [p.name for p in people]
    assert any("Кравченко" in n for n in names)
    assert any("Ерюкова" in n for n in names)


def test_table_row_fio_role_extracted() -> None:
    text = """
Slide 2:
Команда проекта
Кравченко Дмитрий | Тимлид | Backend
Ерюкова Татьяна | Помощник тимлида | Организация
"""
    people = extract_people_from_text(text, section_hint="Команда проекта")
    assert len(people) >= 2


def test_numbered_bullet_team_extracted() -> None:
    text = """
Slide 3:
Участники команды проекта
1. Кравченко Дмитрий — тимлид
2. Ерюкова Татьяна — помощник тимлида
"""
    people = extract_people_from_text(text, section_hint="Участники команды")
    assert len(people) >= 2


def test_pipeline_slide_not_team() -> None:
    text = """
Slide 14:
Посты Telegram
Из Telegram-постов мы строим единый корпус
Qdrant Cloud
Google Colab
"""
    people = extract_people_from_text(text)
    assert len(people) == 0


@pytest.fixture
def grouped_shape_pptx(tmp_path: Path) -> Path:
    path = tmp_path / "grouped.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Команда проекта"
    box2 = group.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
    box2.text_frame.text = "Тимлид: Дмитрий Кравченко"
    prs.save(str(path))
    return path


def test_grouped_shape_text_extracted(grouped_shape_pptx: Path) -> None:
    record = PptxExtractor().extract(grouped_shape_pptx)
    assert "Команда проекта" in record.extracted_text
    assert "Кравченко" in record.extracted_text
    slides = record.metadata.get("slides") or []
    assert slides and slides[0].get("char_count", 0) > 0


def test_ksk_pptx_only_has_team() -> None:
    text = KSK_PRES.read_text(encoding="utf-8")
    contract = _make_builder().build(_pptx_only_extraction(text, "01_presentation.pptx"))
    team = contract.fidelity.team_structured if contract.fidelity else []
    assert len(team) >= 2
    names = [m.name for m in team]
    assert any("Бугров" in n or "Кравченко" in n for n in names)


def test_indlab_pptx_only_no_false_team() -> None:
    text = INDLAB_PRES.read_text(encoding="utf-8")
    contract = _make_builder().build(_pptx_only_extraction(text))
    team = contract.fidelity.team_structured if contract.fidelity else []
    assert len(team) == 0
    exporter = StyledHtmlExporter(ContractRepository(settings))
    html = exporter._render_from_contract(contract, None, ExportTheme.UNIVERSITY_PLATFORM)
    assert "id='team'" not in html and 'id="team"' not in html
    forbidden = ("Посты Telegram", "Из Telegram")
    for name in forbidden:
        assert f"<h3>{name}</h3>" not in html


def test_indlab_multifile_still_has_team() -> None:
    land = (
        Path(__file__).parents[2]
        / "test_corpus"
        / "golden"
        / "indlab_telegram_news"
        / "sources"
        / "02_landing.docx.txt"
    ).read_text(encoding="utf-8")
    pres = INDLAB_PRES.read_text(encoding="utf-8")
    extraction = ExtractionResult(
        project_id=uuid4(),
        payload=ExtractionPayload(),
        files=[
            FileExtraction(
                filename="01_presentation.pptx",
                file_type="pptx",
                extracted_text=pres,
                metadata={"slides_count": pres.count("Slide ")},
            ),
            FileExtraction(
                filename="02_landing.docx",
                file_type="docx",
                extracted_text=land,
            ),
        ],
        extracted_at=utc_now(),
    )
    contract = _make_builder().build(extraction)
    assert contract.fidelity
    assert len(contract.fidelity.team_structured) >= 10
