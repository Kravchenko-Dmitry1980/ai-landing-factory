"""OCR router graceful degradation when engine unavailable."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.config import settings
from app.schemas.extraction import FileExtraction
from app.services.ocr.ocr_router import run_ocr_for_source


@pytest.fixture
def pptx_with_image(tmp_path: Path) -> Path:
    path = tmp_path / "ocr_test.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    try:
        from PIL import Image

        img_path = tmp_path / "x.png"
        Image.new("RGB", (80, 40), color="white").save(img_path)
        slide.shapes.add_picture(str(img_path), Inches(0.5), Inches(0.5))
    except ImportError:
        pytest.skip("Pillow required")
    prs.save(str(path))
    return path


def test_router_no_engine_warning_not_crash(
    pptx_with_image: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", True)

    class _Unavailable:
        name = "paddleocr"

        def is_available(self) -> bool:
            return False

    class _FallbackUnavailable:
        name = "tesseract"

        def is_available(self) -> bool:
            return False

    monkeypatch.setattr(
        "app.services.ocr.ocr_router.PaddleOcrEngine",
        lambda: _Unavailable(),
    )
    monkeypatch.setattr(
        "app.services.ocr.ocr_router.TesseractOcrEngine",
        lambda: _FallbackUnavailable(),
    )

    source = FileExtraction(
        filename="ocr_test.pptx",
        file_type="pptx",
        extracted_text="Slide 1:\n",
        metadata={"source_path": str(pptx_with_image), "slides_count": 1},
    )
    result = run_ocr_for_source(source, missing_fields=["team"])
    assert "OCR engine unavailable." in " ".join(result.warnings)
    assert result.total_chars == 0
