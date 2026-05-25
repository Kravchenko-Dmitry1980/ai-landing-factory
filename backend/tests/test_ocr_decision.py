"""OCR decision unit tests."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.config import settings
from app.schemas.extraction import FileExtraction
from app.services.ocr.ocr_contracts import OcrDecisionAction
from app.services.ocr.ocr_decision import (
    OCR_DISABLED_WARNING,
    should_ocr_source,
)


@pytest.fixture
def image_slide_pptx(tmp_path: Path) -> Path:
    path = tmp_path / "team_image.pptx"
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    try:
        from PIL import Image, ImageDraw

        img_path = tmp_path / "team.png"
        img = Image.new("RGB", (400, 200), color="white")
        draw = ImageDraw.Draw(img)
        draw.text((20, 40), "Команда проекта", fill="black")
        draw.text((20, 80), "Иванов Иван — Тимлид", fill="black")
        draw.text((20, 110), "Петров Петр — Разработчик", fill="black")
        img.save(img_path)
        slide.shapes.add_picture(str(img_path), Inches(0.5), Inches(0.5), width=Inches(8))
    except ImportError:
        pytest.skip("Pillow required for image slide fixture")

    prs.save(str(path))
    return path


def test_pptx_low_text_images_team_missing_should_ocr(
    image_slide_pptx: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", True)
    source = FileExtraction(
        filename="team_image.pptx",
        file_type="pptx",
        extracted_text="Slide 1:\nIntro text only",
        metadata={
            "source_path": str(image_slide_pptx),
            "slides_count": 1,
        },
    )
    decision = should_ocr_source(source, missing_fields=["team"])
    assert decision.action == OcrDecisionAction.RUN
    assert decision.targets
    assert decision.targets[0].source_type == "pptx_slide_image"


def test_pptx_enough_text_skip_ocr(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", True)
    path = tmp_path / "text.pptx"
    prs = Presentation()
    layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Команда проекта"
    slide.placeholders[1].text = (
        "Тимлид: Иванов Иван\n"
        "Помощник тимлида: Петров Петр\n"
        "Участники: Сидоров Сидор, Кузнецова Анна"
    )
    prs.save(str(path))

    source = FileExtraction(
        filename="text.pptx",
        file_type="pptx",
        extracted_text=(
            "Slide 1:\nКоманда проекта\n"
            "Тимлид: Иванов Иван\n"
            "Помощник тимлида: Петров Петр\n"
            "Участники: Сидоров Сидор, Кузнецова Анна"
        ),
        metadata={"source_path": str(path), "slides_count": 1},
    )
    decision = should_ocr_source(source, missing_fields=[])
    assert decision.action == OcrDecisionAction.SKIP


def test_pdf_low_text_should_ocr(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", True)
    source = FileExtraction(
        filename="scan.pdf",
        file_type="pdf",
        extracted_text="",
        metadata={"source_path": "/tmp/scan.pdf", "pages_count": 3},
    )
    decision = should_ocr_source(source, missing_fields=["team"])
    assert decision.action == OcrDecisionAction.RUN
    assert decision.targets
    assert decision.targets[0].source_type == "pdf_page"


def test_ocr_disabled_skipped_with_warning(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "ocr_enabled", False)
    source = FileExtraction(
        filename="x.pptx",
        file_type="pptx",
        extracted_text="",
        metadata={"source_path": "/tmp/x.pptx"},
    )
    decision = should_ocr_source(source, missing_fields=["team"])
    assert decision.action == OcrDecisionAction.DISABLED
    assert OCR_DISABLED_WARNING in decision.warnings[0]
