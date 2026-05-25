"""PPTX image extractor tests."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.services.ocr.renderers.pptx_image_extractor import extract_pptx_images_by_slide


@pytest.fixture
def pptx_with_image_slide(tmp_path: Path) -> Path:
    path = tmp_path / "image_slide.pptx"
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    try:
        from PIL import Image

        img_path = tmp_path / "pic.png"
        Image.new("RGB", (120, 80), color="blue").save(img_path)
        slide.shapes.add_picture(str(img_path), Inches(1), Inches(1), width=Inches(3))
    except ImportError:
        pytest.skip("Pillow required")
    prs.save(str(path))
    return path


def test_extract_image_item_from_image_slide(pptx_with_image_slide: Path) -> None:
    items = extract_pptx_images_by_slide(pptx_with_image_slide)
    assert len(items) >= 1
    assert items[0].slide_index == 1
    assert items[0].image_index == 1
    assert len(items[0].image_bytes) > 0
    assert items[0].ext in ("png", "jpeg", "jpg")
