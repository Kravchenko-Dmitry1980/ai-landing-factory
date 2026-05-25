"""Extract embedded images from PPTX slides."""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

logger = logging.getLogger(__name__)


@dataclass
class PptxImageItem:
    slide_index: int
    image_index: int
    image_bytes: bytes
    ext: str
    width: int | None
    height: int | None
    shape_name: str
    probable_reason: str = "embedded_picture"


def extract_pptx_images_by_slide(path: Path) -> list[PptxImageItem]:
    """Extract all picture shapes grouped by slide index (1-based)."""
    items: list[PptxImageItem] = []
    try:
        prs = Presentation(str(path))
    except Exception as exc:
        logger.warning("PPTX image extraction failed for %s: %s", path, exc)
        return items

    for slide_idx, slide in enumerate(prs.slides, start=1):
        image_index = 0
        for shape in slide.shapes:
            for pic in _iter_picture_shapes(shape):
                try:
                    image = pic.image
                    image_index += 1
                    reason = "team_slide_image" if slide_idx == len(prs.slides) else "embedded_picture"
                    items.append(
                        PptxImageItem(
                            slide_index=slide_idx,
                            image_index=image_index,
                            image_bytes=image.blob,
                            ext=image.ext or "png",
                            width=getattr(pic, "width", None),
                            height=getattr(pic, "height", None),
                            shape_name=getattr(pic, "name", "") or f"picture_{image_index}",
                            probable_reason=reason,
                        )
                    )
                except Exception as pic_exc:
                    logger.debug("Skip picture on slide %s: %s", slide_idx, pic_exc)
    return items


def slide_image_counts(path: Path) -> dict[int, int]:
    """Return number of images per slide index."""
    counts: dict[int, int] = {}
    for item in extract_pptx_images_by_slide(path):
        counts[item.slide_index] = counts.get(item.slide_index, 0) + 1
    return counts


def _iter_picture_shapes(shape):
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        yield shape
        return
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP and hasattr(shape, "shapes"):
        for sub in shape.shapes:
            yield from _iter_picture_shapes(sub)
