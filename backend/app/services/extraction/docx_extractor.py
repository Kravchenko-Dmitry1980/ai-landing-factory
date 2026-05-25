import logging
from pathlib import Path

from docx import Document

from app.schemas.extraction import FileExtraction
from app.services.extraction.base import FileExtractor
from app.services.extraction.utils import file_type_label, normalize_extension

logger = logging.getLogger(__name__)


class DocxExtractor(FileExtractor):
    def supports(self, path: Path) -> bool:
        return normalize_extension(path) == "docx"

    def extract(self, path: Path, filename: str | None = None) -> FileExtraction:
        name = filename or path.name
        warnings: list[str] = []
        errors: list[str] = []

        try:
            doc = Document(path)
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            table_rows: list[str] = []
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        table_rows.append(" | ".join(cells))

            parts = paragraphs + table_rows
            text = "\n\n".join(parts)
            if not text.strip():
                warnings.append("DOCX contains no extractable text")

            metadata = {
                "paragraphs_count": len(paragraphs),
                "tables_count": len(doc.tables),
            }
            return FileExtraction(
                filename=name,
                file_type=file_type_label(path),
                extracted_text=text,
                metadata=metadata,
                warnings=warnings,
                errors=errors,
            )
        except Exception as exc:
            logger.exception("DOCX extraction failed: %s", path)
            return FileExtraction(
                filename=name,
                file_type=file_type_label(path),
                extracted_text="",
                metadata={},
                warnings=["DOCX extraction failed"],
                errors=[str(exc)],
            )
