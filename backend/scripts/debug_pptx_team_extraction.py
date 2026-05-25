#!/usr/bin/env python3
"""Debug PPTX team extraction: slides → markers → candidates → acceptance."""

from __future__ import annotations

import argparse
import sys
import tempfile
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
REPO_ROOT = BACKEND.parent
sys.path.insert(0, str(BACKEND))

from app.services.contract_fidelity.pptx_team_markers import (
    iter_pptx_slides,
    matched_team_markers,
    slide_title_has_team_marker,
    text_has_team_markers,
)
from app.services.evidence.people_extractor import diagnose_people_from_text
from app.services.extraction.dispatcher import ExtractionDispatcher
from app.services.extraction.pptx_extractor import PptxExtractor

PREVIEW_CHARS = 220


def _build_pptx_from_snapshot(snapshot: Path, out_path: Path) -> None:
    import re

    from pptx import Presentation

    text = snapshot.read_text(encoding="utf-8")
    prs = Presentation()
    slides = re.split(r"(?=Slide \d+:)", text)
    for chunk in slides:
        chunk = chunk.strip()
        if not chunk:
            continue
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        lines = [ln.strip() for ln in chunk.splitlines() if ln.strip()]
        title = lines[1] if len(lines) > 1 else (lines[0] if lines else "Slide")
        slide.shapes.title.text = title[:255]
        slide.placeholders[1].text = "\n".join(lines[2:] if len(lines) > 2 else lines)
    if not prs.slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Presentation"
        slide.placeholders[1].text = text[:8000]
    prs.save(out_path)


def _resolve_input(file_path: Path | None, snapshot_path: Path | None) -> tuple[Path, str]:
    if file_path and file_path.is_file():
        return file_path, file_path.name
    if snapshot_path and snapshot_path.is_file():
        tmp = Path(tempfile.mkdtemp(prefix="debug_pptx_"))
        logical = snapshot_path.name.replace(".pptx.txt", ".pptx")
        out = tmp / logical
        _build_pptx_from_snapshot(snapshot_path, out)
        return out, logical
    raise FileNotFoundError("Provide --file or --snapshot")


def _preview(text: str) -> str:
    cleaned = " ".join(text.split())
    if len(cleaned) <= PREVIEW_CHARS:
        return cleaned
    return cleaned[: PREVIEW_CHARS - 1] + "…"


def run_debug(path: Path, logical_name: str) -> int:
    record = PptxExtractor().extract(path, logical_name)
    text = record.extracted_text or ""
    slides_meta = record.metadata.get("slides") or []

    print(f"file={logical_name}")
    print(f"total_slides={record.metadata.get('slides_count', len(slides_meta))}")
    print(f"extracted_chars={len(text)}")
    print(f"global_team_markers={matched_team_markers(text) or 'none'}")
    print()

    all_accepted: list[tuple[int, str, str]] = []
    all_rejected: list[tuple[int, str, str, str]] = []

    parsed_slides = iter_pptx_slides(text)
    for idx, slide_text, title in parsed_slides:
        markers = matched_team_markers(slide_text)
        has_team = slide_title_has_team_marker(title, slide_text) or bool(markers)
        accepted, rejected, _raw = diagnose_people_from_text(
            slide_text,
            section_hint=title or "",
            in_team_section=has_team,
        )
        print(f"--- slide {idx} ---")
        print(f"title: {title or '(none)'}")
        print(f"char_count: {len(slide_text)}")
        print(f"contains_team_markers: {has_team} {markers or ''}")
        print(f"preview: {_preview(slide_text)}")
        if accepted:
            print("accepted:")
            for person in accepted:
                print(f"  + {person.name} ({person.role or 'no role'})")
                all_accepted.append((idx, person.name, person.role or ""))
        if rejected:
            print("rejected:")
            for name, role, reason in rejected:
                print(f"  - {name} ({role or 'no role'}): {reason}")
                all_rejected.append((idx, name, role, reason))
        if not accepted and not rejected and has_team:
            print("  (team markers present but no person candidates)")
        print()

    print(f"final_team_count={len(all_accepted)}")
    if not text_has_team_markers(text):
        print("NOTE: PPTX has no extractable team text markers (тимлид/команда проекта/ФИО).")
    return 0


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        try:
            sys.stdout.reconfigure(encoding="utf-8")
        except Exception:
            pass
    parser = argparse.ArgumentParser(description="Debug PPTX team extraction")
    parser.add_argument("--file", type=Path, default=None)
    parser.add_argument("--snapshot", type=Path, default=None)
    args = parser.parse_args()

    try:
        path, name = _resolve_input(args.file, args.snapshot)
    except FileNotFoundError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    return run_debug(path, name)


if __name__ == "__main__":
    raise SystemExit(main())
