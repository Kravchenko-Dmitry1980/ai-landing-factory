#!/usr/bin/env python3
"""Live HTTP smoke: multi-file upload → contract → team → export."""

from __future__ import annotations

import argparse
import re
import sys
import tempfile
from pathlib import Path

import httpx

BACKEND = Path(__file__).resolve().parents[1]
REPO_ROOT = BACKEND.parent
DEFAULT_CORPUS = REPO_ROOT / "test_corpus" / "golden"
DEFAULT_BACKEND_URL = "http://127.0.0.1:8001"

sys.path.insert(0, str(BACKEND))

from app.schemas.landing_contract import LandingContract  # noqa: E402
from scripts.smoke_corpus import _load_simple_yaml  # noqa: E402
from scripts.smoke_team_export import (  # noqa: E402
    FORBIDDEN_HTML_NAMES,
    _validate_team,
)

FORBIDDEN_TEAM_NAMES = FORBIDDEN_HTML_NAMES + (
    "Google Colab",
    "Qdrant Cloud",
)

STACK_MARKERS = ("Qdrant", "BERTopic", "Neo4j")
TITLE_MARKERS = ("Indlab", "интеллекту", "Telegram", "ассистент", "агрегатор", "новост")


def _build_docx_from_text(text: str, out_path: Path) -> None:
    from docx import Document

    doc = Document()
    for block in re.split(r"\n{2,}", text.strip()):
        paragraph = block.strip()
        if paragraph:
            doc.add_paragraph(paragraph)
    doc.save(out_path)


def _build_pptx_from_text(text: str, out_path: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    slides = re.split(r"(?=Slide \d+:)", text)
    for chunk in slides:
        chunk = chunk.strip()
        if not chunk:
            continue
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        lines = [ln.strip() for ln in chunk.splitlines() if ln.strip()]
        if not lines:
            continue
        title_line = lines[0]
        if title_line.lower().startswith("slide "):
            title_line = lines[1] if len(lines) > 1 else title_line
        slide.shapes.title.text = title_line[:255]
        body = slide.placeholders[1]
        body.text = "\n".join(lines[1:] if lines[0].lower().startswith("slide ") else lines)
    if not prs.slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Presentation"
        slide.placeholders[1].text = text[:8000]
    prs.save(out_path)


def _resolve_corpus_files(project_slug: str) -> list[tuple[str, Path]]:
    sources_dir = DEFAULT_CORPUS / project_slug / "sources"
    if not sources_dir.is_dir():
        raise FileNotFoundError(f"Corpus sources not found: {sources_dir}")

    mapping = [
        ("01_presentation.pptx", "01_presentation.pptx.txt", _build_pptx_from_text),
        ("02_landing.docx", "02_landing.docx.txt", _build_docx_from_text),
    ]
    tmp_dir = Path(tempfile.mkdtemp(prefix="live_multifile_"))
    resolved: list[tuple[str, Path]] = []
    for logical_name, snapshot_name, builder in mapping:
        snapshot = sources_dir / snapshot_name
        if not snapshot.is_file():
            raise FileNotFoundError(f"Missing snapshot: {snapshot}")
        text = snapshot.read_text(encoding="utf-8")
        out = tmp_dir / logical_name
        builder(text, out)
        resolved.append((logical_name, out))
    return resolved


def _prepare_upload_paths(paths: list[Path]) -> list[tuple[str, Path]]:
    prepared: list[tuple[str, Path]] = []
    for path in paths:
        if not path.is_file():
            raise FileNotFoundError(path)
        lower = path.name.lower()
        if lower.endswith(".txt") and ".pptx.txt" in lower:
            tmp = path.with_suffix("").with_suffix(".pptx")
            _build_pptx_from_text(path.read_text(encoding="utf-8"), tmp)
            prepared.append((path.name.replace(".pptx.txt", ".pptx"), tmp))
        elif lower.endswith(".txt") and ".docx.txt" in lower:
            tmp = path.with_suffix("").with_suffix(".docx")
            _build_docx_from_text(path.read_text(encoding="utf-8"), tmp)
            prepared.append((path.name.replace(".docx.txt", ".docx"), tmp))
        else:
            prepared.append((path.name, path))
    return prepared


def _title_ok(title: str) -> bool:
    lower = title.lower()
    return any(marker.lower() in lower for marker in TITLE_MARKERS)


def run_smoke(
    backend_url: str,
    project_name: str,
    upload_files: list[tuple[str, Path]],
    expected: dict | None = None,
) -> list[str]:
    errors: list[str] = []
    base = backend_url.rstrip("/")
    api = f"{base}/api/v1"

    with httpx.Client(timeout=120.0, trust_env=False) as client:
        create_resp = client.post(
            f"{api}/projects",
            json={"name": project_name, "description": "live multifile smoke"},
        )
        if create_resp.status_code != 201:
            errors.append(f"create project HTTP {create_resp.status_code}: {create_resp.text}")
            return errors
        project_id = create_resp.json()["id"]

        multipart = [
            ("files", (name, path.read_bytes(), "application/octet-stream"))
            for name, path in upload_files
        ]
        upload_resp = client.post(f"{api}/projects/{project_id}/upload", files=multipart)
        if upload_resp.status_code != 200:
            errors.append(f"upload HTTP {upload_resp.status_code}: {upload_resp.text}")
            return errors
        uploaded = upload_resp.json().get("files") or []
        if len(uploaded) < len(upload_files):
            errors.append(
                f"upload response files={len(uploaded)} expected>={len(upload_files)}"
            )

        contract_resp = client.get(f"{api}/projects/{project_id}/contract")
        if contract_resp.status_code != 200:
            errors.append(f"contract HTTP {contract_resp.status_code}")
            return errors
        contract = contract_resp.json()
        fidelity = contract.get("fidelity") or {}

        source_count = int(fidelity.get("source_count") or 0)
        if source_count < len(upload_files):
            errors.append(f"source_count={source_count} < uploaded={len(upload_files)}")

        parser_mode = fidelity.get("parser_mode") or ""
        allowed_modes = (
            "field_level_fusion",
            "multi_source_assembly",
            "structured",
            "project_presentation",
        )
        if parser_mode not in allowed_modes:
            errors.append(f"unexpected parser_mode={parser_mode}")
        if len(upload_files) >= 2 and parser_mode != "field_level_fusion":
            errors.append(
                f"multi-file project expected field_level_fusion, got {parser_mode}"
            )

        title = contract.get("title") or ""
        if not _title_ok(title):
            errors.append(f"title does not match Indlab markers: {title!r}")

        gen_resp = client.post(
            f"{api}/projects/{project_id}/generate",
            json={"mode": "stub", "enrich": False},
        )
        if gen_resp.status_code != 200:
            errors.append(f"generate HTTP {gen_resp.status_code}: {gen_resp.text}")

        evidence_resp = client.get(f"{api}/projects/{project_id}/evidence-report")
        if evidence_resp.status_code != 200:
            errors.append(f"evidence-report HTTP {evidence_resp.status_code}")
        else:
            evidence = evidence_resp.json()
            ev_sources = int(evidence.get("source_count") or 0)
            if ev_sources < len(upload_files):
                errors.append(
                    f"evidence source_count={ev_sources} < uploaded={len(upload_files)}"
                )
            if len(upload_files) >= 2:
                field_decisions = evidence.get("field_decisions") or {}
                if "team" not in field_decisions:
                    errors.append("evidence missing field_decisions.team")
                if "modules" not in field_decisions:
                    errors.append("evidence missing field_decisions.modules")

        export_resp = client.get(
            f"{api}/projects/{project_id}/export/html",
            params={"theme": "university_platform"},
        )
        if export_resp.status_code != 200:
            errors.append(f"export HTTP {export_resp.status_code}: {export_resp.text}")
            return errors
        html = export_resp.json().get("html") or ""

        if "theme-university_platform" not in html:
            errors.append("export missing theme-university_platform body class")

        stack_text = html.lower()
        for marker in STACK_MARKERS:
            if marker.lower() not in stack_text:
                errors.append(f"export missing stack marker: {marker}")

        expected = expected or {}
        team_expected = expected.get("must_have_team") or []
        min_team = int(expected.get("min_team") or (10 if team_expected else 1))
        team = fidelity.get("team_structured") or []
        if len(team) < min_team:
            errors.append(f"team_structured={len(team)} < min_team={min_team}")

        contract_model = LandingContract.model_validate(contract)
        team_errors = _validate_team(
            contract_model,
            {
                "min_team": min_team,
                "must_have_team": team_expected,
                "forbidden_team_names": expected.get("forbidden_team_names")
                or list(FORBIDDEN_TEAM_NAMES),
            },
            html,
        )
        errors.extend(team_errors)

        for forbidden in FORBIDDEN_TEAM_NAMES:
            if f"<h3>{forbidden}</h3>" in html:
                errors.append(f"forbidden team-card in export HTML: {forbidden}")

        print(f"project_id={project_id}")
        print(f"uploaded_files={len(uploaded)} source_count={source_count}")
        print(f"parser_mode={parser_mode} team={len(team)}")
        print(f"title={title!r}")

    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description="Live multi-file upload smoke test")
    parser.add_argument("--backend-url", default=DEFAULT_BACKEND_URL)
    parser.add_argument("--project-name", default="Indlab live multifile")
    parser.add_argument("--corpus-project", default="", help="e.g. indlab_telegram_news")
    parser.add_argument("--files", nargs="*", default=[])
    args = parser.parse_args()

    expected: dict | None = None
    if args.corpus_project:
        expected_path = DEFAULT_CORPUS / args.corpus_project / "expected_contract.yml"
        if expected_path.is_file():
            expected = _load_simple_yaml(expected_path)
        upload_files = _resolve_corpus_files(args.corpus_project)
    elif args.files:
        upload_files = _prepare_upload_paths([Path(p) for p in args.files])
    else:
        upload_files = _resolve_corpus_files("indlab_telegram_news")
        expected_path = DEFAULT_CORPUS / "indlab_telegram_news" / "expected_contract.yml"
        if expected_path.is_file():
            expected = _load_simple_yaml(expected_path)

    errors = run_smoke(args.backend_url, args.project_name, upload_files, expected)
    if errors:
        for err in errors:
            print(f"FAIL: {err}", file=sys.stderr)
        return 1
    print("OK: live multifile smoke passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
