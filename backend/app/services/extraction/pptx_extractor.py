import logging
from pathlib import Path

from pptx import Presentation

from app.schemas.extraction import FileExtraction
from app.services.extraction.base import FileExtractor
from app.services.extraction.utils import file_type_label, normalize_extension

logger = logging.getLogger(__name__)


def _shape_text(shape) -> str:
    if not hasattr(shape, "text"):
        return ""
    return (shape.text or "").strip()


class PptxExtractor(FileExtractor):
    def supports(self, path: Path) -> bool:
        return normalize_extension(path) == "pptx"

    def extract(self, path: Path, filename: str | None = None) -> FileExtraction:
        name = filename or path.name
        warnings: list[str] = []
        errors: list[str] = []

        try:
            prs = Presentation(str(path))
            slide_parts: list[str] = []
            notes_parts: list[str] = []
            structured_slides: list[dict] = []
            notes_count = 0

            for idx, slide in enumerate(prs.slides, start=1):
                texts: list[str] = []
                for shape in slide.shapes:
                    t = _shape_text(shape)
                    if t:
                        texts.append(t)
                slide_title = texts[0] if texts else ""
                slide_body = "\n".join(texts[1:]) if len(texts) > 1 else (
                    texts[0] if texts else ""
                )
                note_text = ""
                notes_slide = slide.notes_slide
                if notes_slide is not None:
                    frame = notes_slide.notes_text_frame
                    if frame is not None:
                        note_text = (frame.text or "").strip()
                        if note_text:
                            notes_count += 1
                            notes_parts.append(f"Notes {idx}:\n{note_text}")

                if texts:
                    block = f"Slide {idx}:\n" + "\n".join(texts)
                    slide_parts.append(block)
                    structured_slides.append(
                        {
                            "index": idx,
                            "title": slide_title,
                            "text": slide_body or slide_title,
                            "notes": note_text,
                        }
                    )
                else:
                    warnings.append(f"Slide {idx} has no extractable text")

            body = "\n\n".join(slide_parts)
            if notes_parts:
                body = f"{body}\n\n--- Speaker notes ---\n\n" + "\n\n".join(notes_parts)

            if not body.strip():
                warnings.append("PPTX contains no extractable text")

            metadata = {
                "slides_count": len(prs.slides),
                "notes_slides_count": notes_count,
                "slides": structured_slides,
            }
            return FileExtraction(
                filename=name,
                file_type=file_type_label(path),
                extracted_text=body,
                metadata=metadata,
                warnings=warnings,
                errors=errors,
            )
        except Exception as exc:
            logger.exception("PPTX extraction failed: %s", path)
            return FileExtraction(
                filename=name,
                file_type=file_type_label(path),
                extracted_text="",
                metadata={},
                warnings=["PPTX extraction failed"],
                errors=[str(exc)],
            )
