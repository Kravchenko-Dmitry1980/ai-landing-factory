"""PPTX extractor must retain every slide with text."""

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app.services.extraction.pptx_extractor import PptxExtractor

SLIDE_TEXTS = [
    "Alpha slide title",
    "Beta slide body with unique marker BETA-42",
    "Gamma closing slide GAMMA-99",
]


@pytest.fixture
def synthetic_pptx(tmp_path: Path) -> Path:
    path = tmp_path / "three_slides.pptx"
    prs = Presentation()
    for text in SLIDE_TEXTS:
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = text.split()[0]
        body = slide.placeholders[1]
        body.text = text
    prs.save(str(path))
    return path


def test_pptx_extractor_all_slides(synthetic_pptx: Path) -> None:
    result = PptxExtractor().extract(synthetic_pptx)
    assert "Slide 1:" in result.extracted_text
    assert "Slide 2:" in result.extracted_text
    assert "Slide 3:" in result.extracted_text
    for marker in SLIDE_TEXTS:
        assert marker in result.extracted_text
    assert result.metadata.get("slides_count") == 3
    slides = result.metadata.get("slides") or []
    assert len(slides) >= 3
    assert slides[0]["index"] == 1
    assert slides[2]["index"] == 3
