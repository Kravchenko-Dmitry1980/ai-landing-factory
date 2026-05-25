"""Recursive text extraction from PPTX shapes (groups, tables, text frames)."""

from __future__ import annotations

from pptx.enum.shapes import MSO_SHAPE_TYPE


def extract_shape_texts(shape) -> list[str]:
    """Collect all text fragments from a shape, including nested group/table content."""
    texts: list[str] = []

    try:
        shape_type = shape.shape_type
    except Exception:
        shape_type = None

    if shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            texts.extend(extract_shape_texts(child))
        return texts

    if getattr(shape, "has_table", False):
        table = shape.table
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                texts.append(" | ".join(cells))
        return texts

    if getattr(shape, "has_text_frame", False):
        frame = shape.text_frame
        for paragraph in frame.paragraphs:
            run_text = "".join(run.text for run in paragraph.runs).strip()
            if run_text:
                texts.append(run_text)
            elif paragraph.text.strip():
                texts.append(paragraph.text.strip())
        return texts

    if hasattr(shape, "text"):
        plain = (shape.text or "").strip()
        if plain:
            texts.append(plain)

    return texts


def join_shape_texts(texts: list[str]) -> str:
    """Join shape fragments preserving line breaks between blocks."""
    return "\n".join(t for t in texts if t)
